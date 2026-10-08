import sys, datetime, collections, argparse
from zoneinfo import ZoneInfo
import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("xlsx"); ap.add_argument("head"); ap.add_argument("out")
ap.add_argument("--date")
a = ap.parse_args()
now = datetime.datetime.now(ZoneInfo("America/New_York"))
today = datetime.date.fromisoformat(a.date) if a.date else now.date()
mon = today - datetime.timedelta(days=today.weekday())
sat, sun = mon + datetime.timedelta(5), mon + datetime.timedelta(6)
md = lambda d: f"{d.month}{d.day}"
sl = lambda d: f"{d.month}/{d.day}"
wb = openpyxl.load_workbook(a.xlsx, data_only=True)

def find(base):
    for n in (base, base + " WIP"):
        if n in wb.sheetnames: return n
    return None

def fmt(c):
    if c is None: return ""
    if isinstance(c, datetime.datetime): return c.strftime("%a %-m/%-d").upper()
    if isinstance(c, datetime.time): return c.strftime("%-I:%M %p")
    return str(c).strip().replace("\n", " ")

KIDS = ("kids and", "See Childcare Tab")
def tab_lines(name):
    ws = wb[name]; out = []; drop = None
    for row in ws.iter_rows(values_only=True):
        cells = [fmt(c) for c in row]
        up = [c.upper() for c in cells]
        if "START TIME" in up:
            drop = up.index("CHILDCARE") if "CHILDCARE" in up else None
            tech = up.index("TECH") if "TECH" in up else None
            if isinstance(row[0], datetime.datetime): out.append(f"--- {fmt(row[0])} ---")
            continue
        if isinstance(row[0], datetime.datetime): out.append(f"--- {fmt(row[0])} ---"); continue
        if drop is not None and len(cells) > drop: cells = cells[:drop] + cells[drop + 1:]
        cells = [c for c in cells if not any(k in c for k in KIDS)]
        if any("Combined Childcare" == c or "A2K & Childcare" == c for c in cells): continue
        while cells and cells[-1] == "": cells.pop()
        if any(cells): out.append(" | ".join(cells))
    return out

parts = []
for base, label in ((f"{md(mon)} Mon - {md(sat)} Sat", f"Mon {sl(mon)} to Sat {sl(sat)}"), (f"{md(sun)} Sun", f"Sun {sl(sun)}")):
    n = find(base)
    if not n: parts.append(f"{label}: details are not posted on the Onestop yet."); continue
    if n.endswith("WIP"): parts.append(f"({label} is marked WIP, tentative)")
    parts += tab_lines(n)

ws = wb["\U0001F4C5 "]
ev = collections.defaultdict(list); cur = None
for row in ws.iter_rows(values_only=True):
    cells = (list(row[1:8]) + [None] * 7)[:7]
    if any(isinstance(c, datetime.datetime) for c in cells):
        cur = [c.date() if isinstance(c, datetime.datetime) else None for c in cells]; continue
    if cur is None: continue
    for d, c in zip(cur, cells):
        s = str(c).strip() if c is not None else ""
        if d and s and s.upper() not in ("MON", "TUES", "WED", "THURS", "FRI", "SAT", "SUN"):
            ev[d].append(s.replace("\n", " / "))
first = today.replace(day=1)
cal = [d.strftime("%a %-m/%-d/%Y") + ": " + "; ".join(ev[d]) for d in sorted(ev) if d >= first]
last = max(ev) if ev else today

head = open(a.head).read().replace("__RANGE__", f"Mon {sl(mon)} to Sun {sl(sun)}")
head = head.replace("Oct 2026 through May 2027", f"{first.strftime('%b %Y')} through {last.strftime('%b %Y')}")
stamp = now.strftime("%a %b %-d, %Y at %-I:%M %p") if not a.date else today.strftime("%a %b %-d, %Y")
text = (head + f"ONESTOP COPY (read-only, pulled from the Onestop on {stamp})\n\n"
        f"WEEKLY DETAIL, Mon {sl(mon)} to Sun {sl(sun)}\n"
        "Each line is: GROUP | START | END | WHAT | WHERE | IN CHARGE | HELPERS | FOOD | NOTES. Blank columns at the end are left off.\n\n"
        + "\n".join(parts) + f"\n\nCHURCH CALENDAR, {first.strftime('%b %Y')} to {last.strftime('%b %Y')} (event names only)\n" + "\n".join(cal) + "\n")
for bad in ("—", "–", "Combined Childcare", "kids and"):
    assert bad not in text, bad
assert len(text) > 3000
open(a.out, "w").write(text)
print(len(text), mon, sun)
